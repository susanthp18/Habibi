"""Self-hosted document conversion and chunking for the knowledge base.

AgentStudio replacement for the vendor's hosted /document/process call. It
returns the same shape the ingestion task consumes:

    chunked:        {"docling_metadata": {...}, "chunks": [{chunk_text,
                     chunk_index, contextualized_text, chunk_metadata,
                     token_count}, ...]}
    full_document:  {"docling_metadata": {...}, "full_text": str}

Text is extracted per format (PDF, DOCX, HTML, Markdown, plain text, JSON),
split into heading sections, and paragraphs are packed greedily up to
``max_tokens`` (the hosted service's HybridChunker with merge_peers did the
same). Nothing leaves the box.
"""

from __future__ import annotations

import json
import os
import re

PARSER_NAME = "agentstudio-local"
PARSER_VERSION = 1

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")


def _tokens(text: str) -> int:
    # ponytail: ~4 chars per token; exact counts would need tiktoken, whose
    # encoding file is fetched from the internet on first use. Chunks are far
    # below any embedding model's limit, so the estimate only sizes them.
    return max(1, round(len(text) / 4))


# ---------------------------------------------------------------------------
# Extraction: every format becomes a list of (heading, text, metadata) blocks
# ---------------------------------------------------------------------------


def _markdown_blocks(text: str) -> list[tuple[str | None, str, dict]]:
    blocks: list[tuple[str | None, str, dict]] = []
    heading: str | None = None
    lines: list[str] = []
    for line in text.replace("\r\n", "\n").split("\n"):
        m = _HEADING_RE.match(line)
        if m:
            if any(s.strip() for s in lines):
                blocks.append((heading, "\n".join(lines).strip(), {}))
            heading, lines = m.group(2).strip(), []
            continue
        lines.append(line)
    if any(s.strip() for s in lines):
        blocks.append((heading, "\n".join(lines).strip(), {}))
    return blocks


def _pdf_blocks(path: str) -> tuple[list[tuple[str | None, str, dict]], dict]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    blocks = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append((None, text, {"page": number}))
    if not blocks and reader.pages:
        raise ValueError(
            "No text found in this PDF. Scanned PDFs need OCR, which is not "
            "enabled; upload a text-based PDF or a .docx instead."
        )
    return blocks, {"pages": len(reader.pages)}


def _docx_blocks(path: str) -> list[tuple[str | None, str, dict]]:
    import docx

    document = docx.Document(path)
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "") or ""
        level = style.removeprefix("Heading ").strip()
        if style.startswith("Heading") and level.isdigit():
            lines.append(f"{'#' * min(int(level), 6)} {text}")
        elif style == "Title":
            lines.append(f"# {text}")
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return _markdown_blocks("\n\n".join(lines))


def _html_blocks(path: str) -> list[tuple[str | None, str, dict]]:
    from bs4 import BeautifulSoup

    with open(path, "rb") as fh:
        soup = BeautifulSoup(fh.read(), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    lines: list[str] = []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "pre"]):
        text = " ".join(el.get_text(" ", strip=True).split())
        if not text:
            continue
        if el.name[0] == "h" and el.name[1:].isdigit():
            lines.append(f"{'#' * int(el.name[1:])} {text}")
        else:
            lines.append(text)
    return _markdown_blocks("\n\n".join(lines))


def _read_text(path: str) -> str:
    with open(path, "rb") as fh:
        return fh.read().decode("utf-8", errors="replace")


def extract(path: str, filename: str) -> tuple[list[tuple[str | None, str, dict]], dict]:
    """Return (blocks, metadata) for a supported file, else raise ValueError."""
    ext = os.path.splitext(filename)[1].lower()
    meta: dict = {}
    if ext == ".pdf":
        blocks, meta = _pdf_blocks(path)
    elif ext == ".docx":
        blocks = _docx_blocks(path)
    elif ext in (".html", ".htm"):
        blocks = _html_blocks(path)
    elif ext in (".md", ".markdown", ".txt"):
        blocks = _markdown_blocks(_read_text(path))
    elif ext == ".json":
        blocks = [(None, json.dumps(json.loads(_read_text(path)), indent=2, ensure_ascii=False), {})]
    else:
        raise ValueError(
            f"Unsupported file type '{ext or filename}'. Upload PDF, DOCX, HTML, "
            "Markdown, TXT or JSON."
        )
    return blocks, meta


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------


def _split_long(text: str, max_tokens: int) -> list[str]:
    """Split one oversized paragraph on sentence, then word, boundaries."""
    max_chars = max_tokens * 4
    pieces: list[str] = []
    current = ""
    for sentence in re.split(r"(?<=[.!?])\s+", text):
        words = sentence.split() if len(sentence) > max_chars else [sentence]
        for word in words:
            candidate = f"{current} {word}".strip()
            if current and len(candidate) > max_chars:
                pieces.append(current)
                current = word
            else:
                current = candidate
    if current:
        pieces.append(current)
    return pieces


def chunk_blocks(blocks, max_tokens: int) -> list[dict]:
    """Pack paragraphs within each block up to max_tokens; never cross headings."""
    chunks: list[dict] = []
    for heading, text, meta in blocks:
        paragraphs: list[str] = []
        for para in re.split(r"\n\s*\n", text):
            para = para.strip()
            if not para:
                continue
            if _tokens(para) > max_tokens:
                paragraphs.extend(_split_long(para, max_tokens))
            else:
                paragraphs.append(para)

        current: list[str] = []
        for para in paragraphs + [None]:
            joined = "\n\n".join(current + ([para] if para else []))
            if para is not None and (not current or _tokens(joined) <= max_tokens):
                current.append(para)
                continue
            if current:
                body = "\n\n".join(current)
                chunk_meta = dict(meta)
                if heading:
                    chunk_meta["headings"] = [heading]
                chunks.append(
                    {
                        "chunk_text": body,
                        "chunk_index": len(chunks),
                        "contextualized_text": f"{heading}\n{body}" if heading else body,
                        "chunk_metadata": chunk_meta,
                        "token_count": _tokens(body),
                    }
                )
            current = [para] if para is not None else []
    return chunks


def process_document(
    path: str,
    filename: str,
    *,
    retrieval_mode: str = "chunked",
    max_tokens: int = 128,
) -> dict:
    """Blocking; run it off the event loop."""
    blocks, meta = extract(path, filename)
    metadata = {
        "parser": PARSER_NAME,
        "parser_version": PARSER_VERSION,
        "filename": filename,
        **meta,
    }
    if retrieval_mode == "full_document":
        full_text = "\n\n".join(
            f"## {heading}\n\n{text}" if heading else text for heading, text, _ in blocks
        )
        return {"mode": "full_document", "docling_metadata": metadata, "full_text": full_text}
    chunks = chunk_blocks(blocks, max_tokens)
    metadata["total_chunks"] = len(chunks)
    return {"mode": "chunked", "docling_metadata": metadata, "chunks": chunks}


if __name__ == "__main__":  # self-check: python -m api.services.knowledge_base_local_parser
    sample = [("Fees", "Late fee is 500.\n\nGrace period is 3 days.", {}), (None, "x " * 400, {"page": 2})]
    out = chunk_blocks(sample, max_tokens=128)
    assert out[0]["contextualized_text"].startswith("Fees\n") and "Grace" in out[0]["chunk_text"]
    assert all(c["token_count"] <= 128 for c in out), out
    assert [c["chunk_index"] for c in out] == list(range(len(out)))
    assert out[-1]["chunk_metadata"] == {"page": 2}
    print("ok", len(out), "chunks")
